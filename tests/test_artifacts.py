"""Check the semantic reproducibility of user-facing outputs.

XLSX and PNG carry ZIP/image metadata; comparing them byte-for-byte is
not a meaningful research integrity test. Compare formulas and typed inputs.
"""
from __future__ import annotations

import json
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

from hpf_research.audit import ROOT, records


def test_scenario_configuration_is_explicit():
    config = json.loads((ROOT / "config/analysis_scenarios.json").read_text(encoding="utf-8"))
    assert config["metadata"]["as_of"] == "2025-12-31"
    model = config["loan_interest_comparison"]
    assert model["months"] == 360
    assert model["hpf_rate"] == 0.026
    assert model["commercial_rates"] == [0.031, 0.035, 0.04]
    assert config["liquidity"]["incremental_annual_withdrawal_shocks_pct"] == [20, 50]


def test_loan_and_interest_assumptions_are_workbook_inputs():
    config = json.loads((ROOT / "config/analysis_scenarios.json").read_text(encoding="utf-8"))
    book = load_workbook(ROOT / "outputs/shanghai_hpf_research.xlsx", data_only=False, read_only=True)
    ws = book["LoanSensitivity"]
    model = config["loan_interest_comparison"]
    assert ws.max_row == 1 + len(model["commercial_rates"])
    for offset, commercial in enumerate(model["commercial_rates"], start=2):
        assert ws[f"A{offset}"].value == model["principal_cny"]
        assert abs(ws[f"B{offset}"].value - model["hpf_rate"]) < 1e-10
        assert abs(ws[f"C{offset}"].value - commercial) < 1e-10
        assert ws[f"D{offset}"].value == model["months"]
        for col in "EFGH":
            assert ws[f"{col}{offset}"].data_type == "f"
    book.close()


def test_loan_mix_workbook_references_annual_source_counts():
    book = load_workbook(ROOT / "outputs/shanghai_hpf_research.xlsx", data_only=False, read_only=True)
    annual = book["AnnualRaw"]
    mix = book["LoanMix"]
    assert annual["J11"].value == Decimal("10.90") or annual["J11"].value == 10.90
    assert annual["J12"].value == Decimal("9.42") or annual["J12"].value == 9.42
    for row in range(2, mix.max_row + 1):
        year = int(mix.cell(row, 1).value)
        expected = 11 if year == 2024 else 12
        assert mix.cell(row, 7).value == f"=AnnualRaw!J{expected}*C{row}/100"
    book.close()


def test_fiscal_bridge_workbook_matches_actual_decisions_not_constants():
    book = load_workbook(ROOT / "outputs/shanghai_hpf_research.xlsx", data_only=False, read_only=True)
    ws = book["FiscalBridge"]
    fiscal = {(int(r["year"]), r["basis"], r["metric"]): Decimal(r["value_10k_cny"])
              for r in records("data/raw/center_fiscal_2024_2025.csv")}
    for row in (2, 3):
        year = int(ws.cell(row, 1).value)
        expected = fiscal[(year, "unit_final_accounts", "fiscal_appropriation_income")]
        assert Decimal(str(ws.cell(row, 3).value)) == expected
        assert ws.cell(row, 5).data_type == "f"
    book.close()


def test_loan_adjustments_come_from_official_source_file():
    amendments = records("data/raw/loan_adjustments_2015_2025.csv")
    assert {int(r["year"]) for r in amendments} == {2015, 2016}
    assert all(r["type"] == "securitization_outflow"
               and r["source_url"].startswith("https://www.shzfgjj.cn/")
               for r in amendments)


def test_summary_log_recomputed_from_actual_income_composition():
    log = json.loads((ROOT / "outputs/calculation_log.json").read_text(encoding="utf-8"))
    income = {(int(r["year"]), r["population"]): Decimal(r["high_income_share_pct"])
              for r in records("data/raw/income_composition_2024.csv")}
    observed = income[(2024, "new_loan_borrowers")] / income[(2024, "annual_active_contributors")]
    assert Decimal(log["2024_high_income_loan_vs_contributor_share_ratio"]) == observed


def test_loaded_workbook_has_source_data_and_formula_sheets():
    book = load_workbook(ROOT / "outputs/shanghai_hpf_research.xlsx", data_only=False, read_only=True)
    required = {"AnnualRaw", "DistrictRaw", "DistrictBridge", "CityMonthly", "FiscalBridge",
                "IncomeGroups", "LoanMix", "LoanSensitivity", "LiquiditySensitivity",
                "Sources", "DataGaps"}
    assert required.issubset(book.sheetnames)
    assert book["DistrictRaw"].max_row == 769
    assert book["AnnualRaw"].max_row == 12
    assert book["DistrictBridge"]["L5"].data_type == "f"
    assert book["RepaymentSeasonality"].max_row == 5
    assert book["DistrictSeasonality"].max_row == 65
    assert book["RepaymentSeasonality"]["B5"].value == 16
    book.close()
