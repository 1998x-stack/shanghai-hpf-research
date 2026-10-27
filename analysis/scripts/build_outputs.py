"""Build derived CSVs, a formula-based workbook, figures and calculation log.

Run from the repository root after ``pip install -e '.[analysis]'``:
``python analysis/scripts/build_outputs.py``. Input CSVs and archived originals
are versioned separately; this script never changes them.
"""
from __future__ import annotations

import csv
import json
from decimal import Decimal
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "outputs"
PROCESSED = ROOT / "data/processed"


def read(path: str) -> list[dict[str, str]]:
    with (ROOT / path).open(encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def write(path: Path, rows: list[dict], fields: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fields)
        writer.writeheader()
        writer.writerows(rows)


def d(x: str) -> Decimal:
    return Decimal(x)


def build_data() -> tuple[dict, dict]:
    annual = {int(r["year"]): r for r in read("data/raw/annual_2015_2025.csv")}
    interest = {int(r["year"]): r for r in read("data/raw/interest_expense_2022_2025.csv")}
    monthly: dict[int, list[dict[str, str]]] = {}
    city_rows: list[dict] = []
    bridge_rows: list[dict] = []
    for year in range(2022, 2026):
        rows = read(f"data/processed/district_monthly_{year}_coverage.csv")
        if len(rows) != 192 or len({(r["month"], r["district"]) for r in rows}) != 192:
            raise ValueError(f"{year} district panel is incomplete or duplicated")
        monthly[year] = rows
        previous_collection = previous_withdrawal = Decimal(0)
        for month in range(1, 13):
            group = [r for r in rows if int(r["month"]) == month]
            collection = sum((d(r["year_to_date_collection_10k_cny"]) for r in group), Decimal(0)) / 10000
            withdrawal = sum((d(r["year_to_date_withdrawals_10k_cny"]) for r in group), Decimal(0)) / 10000
            city_rows.append({"year": year, "month": month,
                              "collection_ytd_100m": collection,
                              "withdrawal_ytd_100m": withdrawal,
                              "collection_flow_100m": collection - previous_collection,
                              "withdrawal_flow_100m": withdrawal - previous_withdrawal,
                              "source_scope": "sum of 16 district official month-end YTD rows"})
            previous_collection, previous_withdrawal = collection, withdrawal
        gap = d(annual[year]["contributions_100m"]) - previous_collection
        bridge_rows.append({"year": year, "annual_contribution_100m": annual[year]["contributions_100m"],
                            "district_collection_100m": previous_collection,
                            "gap_100m": gap,
                            "contributor_interest_expense_100m": interest[year]["contributor_interest_expense_100m"],
                            "gap_minus_interest_100m": gap - d(interest[year]["contributor_interest_expense_100m"]),
                            "annual_withdrawal_100m": annual[year]["withdrawals_100m"],
                            "district_withdrawal_100m": previous_withdrawal,
                            "withdrawal_gap_100m": d(annual[year]["withdrawals_100m"]) - previous_withdrawal,
                            "annual_source_url": annual[year]["source_url"],
                            "district_dec_source_url": next(r["source_url"] for r in rows if r["month"] == "12")})
    write(PROCESSED / "city_monthly_flows_2022_2025.csv", city_rows, list(city_rows[0]))
    write(PROCESSED / "annual_district_bridge_2022_2025.csv", bridge_rows, list(bridge_rows[0]))

    composition = read("data/raw/withdrawal_composition_2024_2025.csv")
    composition_derived = []
    for r in composition:
        composition_derived.append({**r, "estimated_amount_100m":
                                    d(r["reported_share_pct"]) * d(r["total_withdrawal_100m"]) / 100,
                                    "estimate_status": "derived from rounded reported share; not original category amount"})
    write(PROCESSED / "withdrawal_composition_estimates.csv", composition_derived,
          list(composition_derived[0]))

    w2024 = {int(r["month"]): d(r["withdrawal_flow_100m"]) for r in city_rows if r["year"] == 2024}
    w2025 = {int(r["month"]): d(r["withdrawal_flow_100m"]) for r in city_rows if r["year"] == 2025}
    annual_growth = sum(w2025.values()) - sum(w2024.values())
    autumn_spring_growth = sum(w2025[m] - w2024[m] for m in (4, 9))
    income = {(int(r["year"]), r["population"]): d(r["high_income_share_pct"])
              for r in read("data/raw/income_composition_2024.csv")}
    high_income_ratio = (income[(2024, "new_loan_borrowers")]
                         / income[(2024, "annual_active_contributors")])
    summary = {
        "district_month_rows_2022_2025": sum(len(v) for v in monthly.values()),
        "2025_withdrawal_growth_100m": str(annual_growth),
        "april_september_share_of_2025_withdrawal_growth_pct": str(100 * autumn_spring_growth / annual_growth),
        "2024_high_income_loan_vs_contributor_share_ratio": str(high_income_ratio),
        "bridge": {str(r["year"]): {k: str(v) for k, v in r.items() if k.endswith("_100m")}
                   for r in bridge_rows},
    }
    OUT.mkdir(exist_ok=True)
    (OUT / "calculation_log.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    return {"annual": annual, "monthly": monthly, "city": city_rows,
            "bridge": bridge_rows, "composition": composition_derived}, summary


def build_workbook(data: dict) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.workbook.properties import CalcProperties

    book = Workbook()
    book.remove(book.active)

    def sheet(name: str, rows: list[dict]) -> None:
        ws = book.create_sheet(name)
        fields = list(rows[0])
        ws.append(fields)
        for row in rows:
            values = []
            for value in row.values():
                if isinstance(value, Decimal):
                    values.append(float(value))
                elif isinstance(value, int):
                    values.append(value)
                elif isinstance(value, str) and value and value.replace(".", "", 1).isdigit():
                    values.append(float(value) if "." in value else int(value))
                else:
                    values.append(value)
            ws.append(values)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
        for i, field in enumerate(fields, 1):
            ws.column_dimensions[get_column_letter(i)].width = min(55, max(13, len(field) + 3))

    annual_rows = [data["annual"][y] for y in sorted(data["annual"])]
    sheet("AnnualRaw", annual_rows)
    ws = book["AnnualRaw"]
    ws.cell(1, 13, "loan_ratio_pct_formula")
    ws.cell(1, 14, "fund_flow_residual_formula")
    ws.cell(1, 15, "loan_flow_residual_formula")
    for i in range(2, 13):
        ws.cell(i, 13, f"=I{i}/F{i}*100")
        if i > 2:
            ws.cell(i, 14, f"=F{i}-F{i-1}-D{i}+E{i}")
            adjustment = "+312.02" if i == 3 else ""
            ws.cell(i, 15, f"=I{i}-I{i-1}-G{i}+H{i}{adjustment}")

    all_monthly = [r for y in sorted(data["monthly"]) for r in data["monthly"][y]]
    sheet("DistrictRaw", all_monthly)
    sheet("CityMonthly", data["city"])
    ws = book["CityMonthly"]
    ws.cell(1, 8, "collection_flow_formula")
    ws.cell(1, 9, "withdrawal_flow_formula")
    for i in range(2, 50):
        jan = ws.cell(i, 2).value == 1
        ws.cell(i, 8, f"=C{i}" if jan else f"=C{i}-C{i-1}")
        ws.cell(i, 9, f"=D{i}" if jan else f"=D{i}-D{i-1}")

    sheet("DistrictBridge", data["bridge"])
    ws = book["DistrictBridge"]
    ws.cell(1, 12, "gap_formula")
    ws.cell(1, 13, "interest_residual_formula")
    ws.cell(1, 14, "withdrawal_gap_formula")
    for i in range(2, 6):
        ws.cell(i, 12, f"=B{i}-C{i}")
        ws.cell(i, 13, f"=L{i}-E{i}")
        ws.cell(i, 14, f"=G{i}-H{i}")

    sheet("Withdrawals", data["composition"])
    ws = book["Withdrawals"]
    ws.cell(1, 9, "amount_formula")
    for i in range(2, 18):
        ws.cell(i, 9, f"=C{i}*D{i}/100")

    for name, path in [
        ("IncomeGroups", "data/raw/income_composition_2024.csv"),
        ("LoanMix", "data/raw/loan_composition_2024_2025.csv"),
        ("ContributorGroups", "data/raw/contributor_composition.csv"),
        ("Financial", "data/raw/financial_items_2024_2025.csv"),
        ("CenterFiscal", "data/raw/center_fiscal_2024_2025.csv"),
        ("ReserveAllocation", "data/raw/reserve_allocation_2015_2025.csv"),
        ("Policy", "policies/policy_timeline.csv"),
        ("MarketContext", "data/external/shanghai_context_2024_2025.csv"),
        ("Sources", "sources/manifest.csv"),
        ("AnnualSourceAudit", "data/processed/annual_source_audit.csv"),
        ("DataGaps", "data/processed/data_gaps.csv"),
    ]:
        sheet(name, read(path))
    ws = book["ReserveAllocation"]
    ws.cell(1, 7, "risk_share_of_hpf_gain_pct_formula")
    for i in range(2, 13):
        if ws.cell(i, 2).value is not None:
            ws.cell(i, 7, f"=D{i}/B{i}*100")
    ws = book["LoanMix"]
    ws.cell(1, 7, "estimated_count_10k_formula")
    for i in range(2, ws.max_row + 1):
        year = int(ws.cell(i, 1).value)
        if year not in data["annual"]:
            raise ValueError(f"LoanMix unknown source year {year}")
        annual_row = sorted(data["annual"]).index(year) + 2
        ws.cell(i, 7, f"=AnnualRaw!J{annual_row}*C{i}/100")

    center = {(int(r["year"]), r["basis"], r["metric"]): d(r["value_10k_cny"])
              for r in read("data/raw/center_fiscal_2024_2025.csv")}
    management = {(int(r["year"]), r["metric"]): d(r["amount_100m"])
                  for r in read("data/raw/financial_items_2024_2025.csv")}
    fiscal = [
        {"year": year, "reported_management_fee_100m": management[(year, "management_fee_assigned")],
         "actual_fiscal_appropriation_10k": center[(year, "unit_final_accounts",
                                                      "fiscal_appropriation_income")],
         "rounded_fee_10k_formula": "", "difference_10k_formula": "",
         "rounding_half_unit_10k": 50,
         "status": "amounts compatible with common transfer route; transactions not reconciled"}
        for year in (2024, 2025)
    ]
    sheet("FiscalBridge", fiscal)
    ws = book["FiscalBridge"]
    for i in range(2, 4):
        ws.cell(i, 4, f"=B{i}*10000")
        ws.cell(i, 5, f"=C{i}-D{i}")

    config = json.loads((ROOT / "config/analysis_scenarios.json").read_text(encoding="utf-8"))
    model = config["loan_interest_comparison"]
    scenarios = []
    for commercial_rate in model["commercial_rates"]:
        scenarios.append({"principal_cny": model["principal_cny"], "hpf_rate": model["hpf_rate"],
                          "commercial_rate": commercial_rate, "months": model["months"],
                          "hpf_payment_formula": "", "commercial_payment_formula": "",
                          "nominal_saving_formula": "", "present_value_formula": "",
                          "status": "hypothetical matched-risk commercial rate; no observed subsidy total"})
    sheet("LoanSensitivity", scenarios)
    ws = book["LoanSensitivity"]
    for i in range(2, 5):
        ws.cell(i, 5, f"=PMT(B{i}/12,D{i},-A{i})")
        ws.cell(i, 6, f"=PMT(C{i}/12,D{i},-A{i})")
        ws.cell(i, 7, f"=(F{i}-E{i})*D{i}")
        ws.cell(i, 8, f"=(F{i}-E{i})*(1-(1+C{i}/12)^(-D{i}))/(C{i}/12)")

    liquidity = [
        {"annual_withdrawals_100m": d(data["annual"][2025]["withdrawals_100m"]),
         "shock_pct": shock,
         "year_end_deposits_100m": management[(2025, "deposit_balance")],
         "incremental_withdrawal_formula": "",
         "share_of_year_end_deposits_formula": "",
         "limitation": "Illustrative annual shock; not a 1-90 day liquidity test"}
        for shock in config["liquidity"]["incremental_annual_withdrawal_shocks_pct"]
    ]
    sheet("LiquiditySensitivity", liquidity)
    ws = book["LiquiditySensitivity"]
    for i in range(2, 4):
        ws.cell(i, 4, f"=A{i}*B{i}/100")
        ws.cell(i, 5, f"=D{i}/C{i}*100")
    book.calculation = CalcProperties(calcMode="auto", fullCalcOnLoad=True)
    book.save(OUT / "shanghai_hpf_research.xlsx")


def build_figures(data: dict) -> None:
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    figdir = OUT / "figures"
    figdir.mkdir(exist_ok=True)
    annual = data["annual"]
    years = list(sorted(annual))
    ratios = [100 * float(annual[y]["loan_balance_100m"]) / float(annual[y]["fund_balance_100m"])
              for y in years]
    fig, ax = plt.subplots(figsize=(8, 4.2))
    ax.plot(years, ratios, marker="o", color="#174b78")
    ax.set(ylabel="Loan balance / fund balance (%)", xlabel="Year",
           title="Shanghai HPF personal loan ratio, 2015–2025")
    ax.grid(alpha=.25)
    fig.tight_layout()
    fig.savefig(figdir / "loan_ratio_2015_2025.png", dpi=180)
    plt.close(fig)

    fig, ax = plt.subplots(figsize=(9, 4.6))
    for year, color in [(2023, "#969696"), (2024, "#4b83b7"), (2025, "#dc6a33")]:
        rows = [r for r in data["city"] if r["year"] == year]
        ax.plot([r["month"] for r in rows], [float(r["withdrawal_flow_100m"]) for r in rows],
                marker="o", color=color, label=str(year))
    ax.set(xlabel="Month", ylabel="Monthly withdrawals (CNY 100m)",
           title="Shanghai HPF withdrawals, district YTD differences")
    ax.set_xticks(range(1, 13))
    ax.grid(alpha=.25)
    ax.legend()
    fig.tight_layout()
    fig.savefig(figdir / "withdrawal_seasonality_2023_2025.png", dpi=180)
    plt.close(fig)


def main() -> None:
    data, summary = build_data()
    build_workbook(data)
    build_figures(data)
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
